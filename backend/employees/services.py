"""Employee bulk import from an Excel or CSV upload."""

from __future__ import annotations

import csv
import io
from decimal import Decimal, InvalidOperation

import openpyxl

from companies.models import Branch
from employees.models import Employee

#: Accepted spreadsheet headers -> model field. Arabic first, English accepted.
HEADER_MAP = {
    "الكود": "code",
    "code": "code",
    "الاسم": "name_ar",
    "name": "name_ar",
    "name_ar": "name_ar",
    "الاسم بالانجليزية": "name_en",
    "name_en": "name_en",
    "الفرع": "branch",
    "branch": "branch",
    "الوظيفة": "job_title",
    "job_title": "job_title",
    "ساعات العمل": "daily_hours",
    "daily_hours": "daily_hours",
    "الراتب": "base_salary",
    "base_salary": "base_salary",
    "الراتب التاميني": "insurable_salary",
    "insurable_salary": "insurable_salary",
}

TEMPLATE_HEADERS = [
    "الكود",
    "الاسم",
    "الفرع",
    "الوظيفة",
    "ساعات العمل",
    "الراتب",
    "الراتب التاميني",
]


def _normalise(value) -> str:
    return str(value or "").strip()


def _decimal(value, default="0") -> Decimal:
    raw = _normalise(value).replace(",", "")
    if not raw:
        return Decimal(default)
    try:
        return Decimal(raw)
    except InvalidOperation as exc:
        raise ValueError(f"'{value}' is not a number") from exc


def read_upload(uploaded) -> list[dict]:
    """Parse an .xlsx or .csv upload into header-mapped dicts."""
    name = (getattr(uploaded, "name", "") or "").lower()
    if name.endswith(".csv"):
        return _read_csv(uploaded)
    return _read_xlsx(uploaded)


def _read_csv(uploaded) -> list[dict]:
    text = uploaded.read().decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return []
    return _map_rows(rows[0], rows[1:])


def _read_xlsx(uploaded) -> list[dict]:
    workbook = openpyxl.load_workbook(uploaded, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    if not rows:
        return []
    return _map_rows(rows[0], rows[1:])


def _map_rows(header_row, data_rows) -> list[dict]:
    columns = {}
    for index, raw in enumerate(header_row):
        key = HEADER_MAP.get(_normalise(raw))
        if key:
            columns[key] = index

    parsed = []
    for offset, raw_row in enumerate(data_rows, start=2):
        values = {
            field: raw_row[index] if index < len(raw_row) else None
            for field, index in columns.items()
        }
        if not _normalise(values.get("code")) and not _normalise(values.get("name_ar")):
            continue
        values["_row"] = offset
        parsed.append(values)
    return parsed


def validate_rows(company, rows: list[dict]) -> tuple[list[dict], list[dict]]:
    """Split parsed rows into (valid, errors), resolving branches by name."""
    branches = {branch.name_ar: branch for branch in Branch.objects.for_company(company)}
    existing_codes = set(Employee.objects.for_company(company).values_list("code", flat=True))

    valid, errors = [], []
    seen: set[str] = set()

    for row in rows:
        line = row.get("_row")
        code = _normalise(row.get("code"))
        name = _normalise(row.get("name_ar"))
        branch_name = _normalise(row.get("branch"))
        problems = []

        if not code:
            problems.append("Missing code")
        elif code in seen:
            problems.append("Duplicate code in this file")
        if not name:
            problems.append("Missing name")
        if branch_name and branch_name not in branches:
            problems.append(f"Unknown branch '{branch_name}'")
        if not branch_name:
            problems.append("Missing branch")

        try:
            base_salary = _decimal(row.get("base_salary"))
        except ValueError as exc:
            problems.append(f"Salary: {exc}")
            base_salary = Decimal("0")

        try:
            insurable = _decimal(row.get("insurable_salary"))
        except ValueError as exc:
            problems.append(f"Insurable salary: {exc}")
            insurable = Decimal("0")

        try:
            hours = _decimal(row.get("daily_hours"), default="9")
        except ValueError as exc:
            problems.append(f"Daily hours: {exc}")
            hours = Decimal("9")

        if problems:
            errors.append({"row": line, "code": code, "name": name, "errors": problems})
            continue

        seen.add(code)
        valid.append(
            {
                "row": line,
                "code": code,
                "name_ar": name,
                "name_en": _normalise(row.get("name_en")),
                "branch": branches[branch_name],
                "branch_name": branch_name,
                "job_title": _normalise(row.get("job_title")),
                "daily_hours": hours,
                "base_salary": base_salary,
                "insurable_salary": insurable,
                "action": "update" if code in existing_codes else "create",
            }
        )

    return valid, errors


def commit_rows(company, rows: list[dict]) -> dict:
    """Create or update employees from already-validated rows."""
    created = updated = 0
    for row in rows:
        _, was_created = Employee.all_objects.update_or_create(
            company=company,
            code=row["code"],
            defaults={
                "name_ar": row["name_ar"],
                "name_en": row.get("name_en", ""),
                "branch": row["branch"],
                "job_title": row.get("job_title", ""),
                "daily_hours": row["daily_hours"],
                "base_salary": row["base_salary"],
                "insurable_salary": row["insurable_salary"],
                "is_active": True,
            },
        )
        created += int(was_created)
        updated += int(not was_created)
    return {"created": created, "updated": updated}


def build_template() -> bytes:
    """An empty import workbook with the expected Arabic headers."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "الموظفين"
    sheet.sheet_view.rightToLeft = True
    sheet.append(TEMPLATE_HEADERS)
    for index, _header in enumerate(TEMPLATE_HEADERS, start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = 18

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
