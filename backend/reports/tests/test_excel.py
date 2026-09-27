"""The generated workbook must be a working spreadsheet, not a picture of one.

These tests re-open the file openpyxl produced, read the *formulas* out of it,
evaluate them against the sheet's own input cells, and compare the result with
the `PayrollLine` rows in the database. If someone changes a formula template
and breaks the arithmetic, this fails even though the engine still agrees with
itself.
"""

import io
from decimal import Decimal

import openpyxl
import pytest

from core.factories import (
    BranchFactory,
    CompanyFactory,
    EmployeeFactory,
    PayrollPolicyFactory,
    PeriodFactory,
)
from payroll.models import PayrollLine
from payroll.services.runner import run_period
from reports.services import excel
from reports.services.formula_eval import evaluate
from reports.services.generator import generate_excel, lines_for
from reports.services.layout import (
    COLUMNS,
    FIRST_DATA_ROW,
    FORMULA_FIELD,
    MASTER_SHEET_AR,
    STRIP_PREFIX_AR,
)

pytestmark = pytest.mark.django_db

TOLERANCE = Decimal("0.01")
D = Decimal


@pytest.fixture
def seeded_period():
    """A small company whose numbers exercise every calculated column."""
    company = CompanyFactory(slug="excel-co")
    PayrollPolicyFactory(company=company)
    branch = BranchFactory(company=company, name_ar="فرع الاختبار")

    for index in range(3):
        EmployeeFactory(
            company=company,
            branch=branch,
            code=f"E{index + 1}",
            base_salary=D("6000") + D(index) * D("500"),
            insurable_salary=D("2000"),
            daily_hours=D("9"),
        )

    period = PeriodFactory(company=company)
    run_period(period)

    # Give every calculated column something to chew on.
    for line in PayrollLine.all_objects.filter(company=company, period=period):
        line.work_days = D("26")
        line.overtime_hours = D("12")
        line.leave_allowance_days = D("2")
        line.bonus = D("300")
        line.other_earnings = D("150")
        line.late_hours = D("5")
        line.advances = D("400")
        line.carried_advance = D("100")
        line.admin_penalty_days = D("1")
        line.fingerprint_penalty_days = D("2")
        line.unexcused_absence_days = D("1")
        line.sick_days = D("3")
        line.deviations = D("50")
        line.shortage_custody = D("75")
        line.overrides = {"work_days": "0.00"}
        line.save()

    run_period(period)
    return period


def _cells(sheet):
    """Accessor the formula evaluator uses: (row, 'AB') -> value."""

    def read(row: int, column: str):
        return sheet[f"{column}{row}"].value

    return read


def _open(content: bytes, data_only: bool = False):
    return openpyxl.load_workbook(io.BytesIO(content), data_only=data_only)


def test_workbook_has_a_sheet_per_branch_and_payslip_strips(seeded_period):
    content = excel.generate_workbook(seeded_period, lines_for(seeded_period))
    workbook = _open(content)

    assert MASTER_SHEET_AR in workbook.sheetnames
    assert "فرع الاختبار" in workbook.sheetnames
    assert f"{STRIP_PREFIX_AR} فرع الاختبار" in workbook.sheetnames


def test_every_sheet_is_right_to_left(seeded_period):
    workbook = _open(excel.generate_workbook(seeded_period, lines_for(seeded_period)))

    for sheet in workbook.worksheets:
        assert sheet.sheet_view.rightToLeft is True, sheet.title


def test_title_and_two_row_header_match_the_workbook_layout(seeded_period):
    workbook = _open(excel.generate_workbook(seeded_period, lines_for(seeded_period)))
    sheet = workbook[MASTER_SHEET_AR]

    assert sheet["A1"].value.startswith(MASTER_SHEET_AR)
    merges = {str(cell_range) for cell_range in sheet.merged_cells.ranges}
    assert "A1:AG1" in merges
    assert "G2:O2" in merges  # الاستحقاق
    assert "P2:AE2" in merges  # الاستقطاع
    assert "A2:A3" in merges  # الكود spans both header rows

    assert sheet["A2"].value == "الكود"
    assert sheet["G2"].value == "الاستحقاق"
    assert sheet["P2"].value == "الاستقطاع"
    assert sheet["AG2"].value == "صافى الراتب"


def test_calculated_columns_are_written_as_live_formulas(seeded_period):
    workbook = _open(excel.generate_workbook(seeded_period, lines_for(seeded_period)))
    sheet = workbook[MASTER_SHEET_AR]

    assert sheet[f"H{FIRST_DATA_ROW}"].value == f"=+F{FIRST_DATA_ROW}/30*G{FIRST_DATA_ROW}"
    assert (
        sheet[f"W{FIRST_DATA_ROW}"].value == f"=+H{FIRST_DATA_ROW}/30*V{FIRST_DATA_ROW}"
    ), "fingerprint penalty must divide the earned value H, not the salary F"
    assert sheet[f"AC{FIRST_DATA_ROW}"].value == f"=AB{FIRST_DATA_ROW}*11%"
    assert sheet[f"AG{FIRST_DATA_ROW}"].value == f"=+O{FIRST_DATA_ROW}-AF{FIRST_DATA_ROW}"


def test_formulas_in_the_file_reproduce_every_stored_value(seeded_period):
    """The verification that matters: evaluate the file's own formulas."""
    lines = lines_for(seeded_period)
    workbook = _open(excel.generate_workbook(seeded_period, lines))
    sheet = workbook[MASTER_SHEET_AR]
    read = _cells(sheet)

    failures = []
    for offset, line in enumerate(lines):
        row = FIRST_DATA_ROW + offset
        for letter, field in FORMULA_FIELD.items():
            formula = sheet[f"{letter}{row}"].value
            computed = evaluate(str(formula), read)
            stored = Decimal(getattr(line, field))
            if abs(computed - stored) > TOLERANCE:
                failures.append(
                    f"row {row} {letter} ({field}): formula={formula} -> {computed}, db={stored}"
                )

    assert not failures, "\n".join(failures)


def test_totals_rows_sum_the_columns_they_claim_to(seeded_period):
    lines = lines_for(seeded_period)
    workbook = _open(excel.generate_workbook(seeded_period, lines))
    sheet = workbook[MASTER_SHEET_AR]
    read = _cells(sheet)

    totals_row = FIRST_DATA_ROW + len(lines)
    subtotal_row = totals_row + 1

    assert sheet[f"A{totals_row}"].value == "الاجماليات"
    assert sheet[f"A{subtotal_row}"].value == "الإجمالي"

    expected_net = sum((Decimal(line.net_salary) for line in lines), Decimal(0))
    assert abs(evaluate(str(sheet[f"AG{totals_row}"].value), read) - expected_net) <= TOLERANCE

    expected_base = sum((Decimal(line.base_salary) for line in lines), Decimal(0))
    assert abs(evaluate(str(sheet[f"F{totals_row}"].value), read) - expected_base) <= TOLERANCE

    # AD had no SUM in the source workbook; ours does.
    assert sheet[f"AD{totals_row}"].value.startswith("=SUM(")


def test_payslip_strips_use_four_rows_per_employee(seeded_period):
    lines = lines_for(seeded_period)
    workbook = _open(excel.generate_workbook(seeded_period, lines))
    sheet = workbook[f"{STRIP_PREFIX_AR} فرع الاختبار"]

    for index, line in enumerate(lines):
        data_row = 2 + index * 4 + 2
        assert sheet[f"A{data_row}"].value == line.employee_code
        assert sheet[f"B{data_row}"].value == line.employee_name
        # The row after each strip is the blank separator you cut along.
        assert sheet[f"A{data_row + 1}"].value is None


def test_policy_values_are_written_into_the_formulas(seeded_period):
    """A company with different rules gets a workbook that says so."""
    policy = seeded_period.company.payrollpolicys.first()
    policy.sick_deduction_rate = Decimal("0.5")
    policy.insurance_employee_rate = Decimal("0.14")
    policy.month_day_basis = 26
    policy.overtime_multiplier = Decimal("1.5")
    policy.save()

    workbook = _open(excel.generate_workbook(seeded_period, lines_for(seeded_period), policy))
    sheet = workbook[MASTER_SHEET_AR]
    row = FIRST_DATA_ROW

    assert sheet[f"AA{row}"].value == f"=+(F{row}/26*Z{row})*0.5"
    assert sheet[f"AC{row}"].value == f"=AB{row}*14%"
    assert sheet[f"H{row}"].value == f"=+F{row}/26*G{row}"
    assert sheet[f"J{row}"].value == f"=+I{row}*F{row}/26/E{row}*1.5"


def test_generate_excel_stores_a_report_file(seeded_period):
    report = generate_excel(seeded_period)

    assert report.pk is not None
    assert report.kind == "xlsx"
    assert report.file.name.endswith(".xlsx")
    assert report.company_id == seeded_period.company_id

    report.file.open("rb")
    try:
        workbook = _open(report.file.read())
    finally:
        report.file.close()
    assert MASTER_SHEET_AR in workbook.sheetnames


def test_number_format_is_two_decimals(seeded_period):
    workbook = _open(excel.generate_workbook(seeded_period, lines_for(seeded_period)))
    sheet = workbook[MASTER_SHEET_AR]

    assert sheet[f"F{FIRST_DATA_ROW}"].number_format == "#,##0.00"
    assert sheet[f"AG{FIRST_DATA_ROW}"].number_format == "#,##0.00"


def test_english_report_language_produces_ltr_sheets(seeded_period):
    workbook = _open(excel.generate_workbook(seeded_period, lines_for(seeded_period), None, "en"))

    sheet = workbook["Payroll report"]
    assert sheet.sheet_view.rightToLeft is False
    assert sheet["A2"].value == "Code"
    assert sheet["G2"].value == "Earnings"


def test_every_column_of_the_layout_is_written(seeded_period):
    """Guard against a column being dropped from the layout table."""
    workbook = _open(excel.generate_workbook(seeded_period, lines_for(seeded_period)))
    sheet = workbook[MASTER_SHEET_AR]

    for column in COLUMNS:
        header_cell = (
            sheet[f"{column.letter}2"] if column.group is None else sheet[f"{column.letter}3"]
        )
        assert header_cell.value == column.header_ar, column.letter
