"""Excel report generator.

Reproduces `docs/reference/accounting_monthly.xlsx`: right-to-left sheets, the
two-row Arabic header with its group merges, one sheet per branch, payslip-strip
sheets, and — crucially — **live formulas** rather than baked values, so the
downloaded file keeps working like the original spreadsheet.

Policy values that differ from the workbook's hard-coded constants are written
into the formulas themselves (a 30% sick rate emits `*0.3`), so the file always
explains its own arithmetic.
"""

from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from companies.models import FingerprintPenaltyBase
from reports.services.layout import (
    ARABIC_MONTHS,
    COLUMNS,
    FIRST_DATA_ROW,
    FIRST_DEDUCTION_COLUMN,
    FIRST_GROUP_COLUMN,
    GROUP_DEDUCTIONS_AR,
    GROUP_DEDUCTIONS_EN,
    GROUP_EARNINGS_AR,
    GROUP_EARNINGS_EN,
    HEADER_SUB_ROW,
    HEADER_TOP_ROW,
    LAST_DEDUCTION_COLUMN,
    LAST_EARNINGS_COLUMN,
    LAST_TITLE_COLUMN,
    MASTER_SHEET_AR,
    STRIP_PREFIX_AR,
    SUBTOTAL_LABEL_AR,
    SUBTOTAL_LABEL_EN,
    TITLE_ROW,
    TOTALS_LABEL_AR,
    TOTALS_LABEL_EN,
    Column,
)

MONEY_FORMAT = "#,##0.00"
QUANTITY_FORMAT = "#,##0.##"

HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
TITLE_FILL = PatternFill("solid", fgColor="BDD7EE")
TOTALS_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="8EA9DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEADER_FONT = Font(bold=True, size=10)
TITLE_FONT = Font(bold=True, size=14)
TOTALS_FONT = Font(bold=True, size=10)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _decimal_text(value: Decimal) -> str:
    """Render a policy constant the way a person would type it (0.25, not 0.2500)."""
    text = format(Decimal(value).normalize(), "f")
    return text


class PolicyFormulas:
    """Turns a PayrollPolicy into the literals embedded in each formula."""

    def __init__(self, policy=None):
        self.basis = int(policy.month_day_basis) if policy else 30
        self.overtime = Decimal(policy.overtime_multiplier) if policy else Decimal(1)
        self.absence = Decimal(policy.absence_multiplier) if policy else Decimal(1)
        self.sick = Decimal(policy.sick_deduction_rate) if policy else Decimal("0.25")
        self.insurance = Decimal(policy.insurance_employee_rate) if policy else Decimal("0.11")
        self.fingerprint_base = (
            policy.fingerprint_penalty_base if policy else FingerprintPenaltyBase.EARNED
        )

    @property
    def tokens(self) -> dict[str, str]:
        return {
            "basis": str(self.basis),
            # The workbook writes no multiplier at all when it is 1.
            "ot": "" if self.overtime == 1 else f"*{_decimal_text(self.overtime)}",
            "absence": _decimal_text(self.absence),
            "sick": _decimal_text(self.sick),
            "insurance": _decimal_text(self.insurance * 100),
            "fp": "F" if self.fingerprint_base == FingerprintPenaltyBase.BASE else "H",
        }

    def formula_for(self, column: Column, row: int) -> str:
        return column.formula.format(r=row, **self.tokens)


def _title_for(period, branch_name: str | None, language: str) -> str:
    if language == "en":
        month = period.start_date.strftime("%B")
        if branch_name:
            return f"{branch_name} payroll — {month} {period.year}"
        return f"Payroll report for {month} {period.year}"
    if branch_name:
        return f"مرتبات {branch_name} عن شهر {period.month}-{period.year}"
    return f"{MASTER_SHEET_AR} عن شهر {ARABIC_MONTHS[period.month]} {period.year}"


def _header(column: Column, language: str) -> str:
    return column.header_en if language == "en" else column.header_ar


def _write_header(sheet: Worksheet, period, branch_name: str | None, language: str) -> None:
    """Title row plus the two-row header with its group merges."""
    rtl = language != "en"
    sheet.sheet_view.rightToLeft = rtl

    sheet.merge_cells(f"A{TITLE_ROW}:{LAST_TITLE_COLUMN}{TITLE_ROW}")
    title = sheet.cell(row=TITLE_ROW, column=1, value=_title_for(period, branch_name, language))
    title.font = TITLE_FONT
    title.alignment = CENTER
    title.fill = TITLE_FILL
    sheet.row_dimensions[TITLE_ROW].height = 26

    earnings = GROUP_EARNINGS_EN if language == "en" else GROUP_EARNINGS_AR
    deductions = GROUP_DEDUCTIONS_EN if language == "en" else GROUP_DEDUCTIONS_AR
    sheet.merge_cells(
        f"{FIRST_GROUP_COLUMN}{HEADER_TOP_ROW}:{LAST_EARNINGS_COLUMN}{HEADER_TOP_ROW}"
    )
    sheet.merge_cells(
        f"{FIRST_DEDUCTION_COLUMN}{HEADER_TOP_ROW}:{LAST_DEDUCTION_COLUMN}{HEADER_TOP_ROW}"
    )
    for letter, label in (
        (FIRST_GROUP_COLUMN, earnings),
        (FIRST_DEDUCTION_COLUMN, deductions),
    ):
        cell = sheet[f"{letter}{HEADER_TOP_ROW}"]
        cell.value = label
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.fill = HEADER_FILL

    for column in COLUMNS:
        label = _header(column, language)
        if column.group is None:
            # A..F and AF..AH span both header rows.
            sheet.merge_cells(f"{column.letter}{HEADER_TOP_ROW}:{column.letter}{HEADER_SUB_ROW}")
            cell = sheet[f"{column.letter}{HEADER_TOP_ROW}"]
        else:
            cell = sheet[f"{column.letter}{HEADER_SUB_ROW}"]
        cell.value = label
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.fill = HEADER_FILL

    for row in (HEADER_TOP_ROW, HEADER_SUB_ROW):
        sheet.row_dimensions[row].height = 30
        for column in COLUMNS:
            sheet[f"{column.letter}{row}"].border = BORDER

    for column in COLUMNS:
        sheet.column_dimensions[column.letter].width = column.width


def _write_line(sheet: Worksheet, row: int, line, policy: PolicyFormulas) -> None:
    """One employee row: inputs as values, calculated columns as live formulas."""
    hours_are_zero = Decimal(line.daily_hours or 0) <= 0

    for column in COLUMNS:
        cell = sheet.cell(row=row, column=column.index)
        cell.border = BORDER

        if column.kind == "formula":
            # An employee with zero daily hours would make F/30/E a #DIV/0!.
            if hours_are_zero and column.letter in {"J", "Q"}:
                cell.value = 0
            else:
                cell.value = policy.formula_for(column, row)
            cell.number_format = MONEY_FORMAT
            continue

        value = getattr(line, column.field) if column.field else None
        if column.kind == "text":
            cell.value = value or ""
            cell.alignment = Alignment(horizontal="right" if column.letter != "A" else "center")
        else:
            cell.value = float(value or 0)
            cell.number_format = MONEY_FORMAT if column.kind == "money" else QUANTITY_FORMAT


def _write_totals(
    sheet: Worksheet, first_row: int, last_row: int, language: str
) -> tuple[int, int]:
    """`الاجماليات` (SUM) then `الإجمالي` (SUBTOTAL 9), as in the workbook."""
    totals_row = last_row + 1
    subtotal_row = totals_row + 1

    labels = (
        (TOTALS_LABEL_EN, SUBTOTAL_LABEL_EN)
        if language == "en"
        else (TOTALS_LABEL_AR, SUBTOTAL_LABEL_AR)
    )

    for row, label in zip((totals_row, subtotal_row), labels, strict=True):
        cell = sheet.cell(row=row, column=1, value=label)
        cell.font = TOTALS_FONT
        cell.fill = TOTALS_FILL
        cell.alignment = CENTER

    for column in COLUMNS:
        if not column.is_numeric:
            continue
        span = f"{column.letter}{first_row}:{column.letter}{last_row}"

        total = sheet.cell(row=totals_row, column=column.index)
        if column.letter == "O":
            # The workbook recomposes O from its components rather than summing.
            total.value = f"=+N{totals_row}+M{totals_row}+L{totals_row}+J{totals_row}+H{totals_row}"
        else:
            total.value = f"=SUM({span})"

        subtotal = sheet.cell(row=subtotal_row, column=column.index)
        if column.letter == "AG":
            subtotal.value = f"=+O{subtotal_row}-AF{subtotal_row}"
        else:
            subtotal.value = f"=SUBTOTAL(9,{span})"

        for cell in (total, subtotal):
            cell.number_format = MONEY_FORMAT
            cell.font = TOTALS_FONT
            cell.fill = TOTALS_FILL
            cell.border = BORDER

    # AH on the subtotal row: base salary total minus net total.
    remainder = sheet.cell(row=subtotal_row, column=34)
    remainder.value = f"=SUM(F{subtotal_row}-AG{subtotal_row})"
    remainder.number_format = MONEY_FORMAT
    remainder.font = TOTALS_FONT
    remainder.fill = TOTALS_FILL

    return totals_row, subtotal_row


def _configure_print(sheet: Worksheet) -> None:
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = f"{HEADER_TOP_ROW}:{HEADER_SUB_ROW}"


def build_sheet(
    workbook: Workbook,
    title: str,
    period,
    lines,
    policy,
    language: str,
    branch_name: str | None = None,
) -> Worksheet:
    """A full report sheet: header, one row per line, both totals rows."""
    sheet = workbook.create_sheet(title=title[:31])
    _write_header(sheet, period, branch_name, language)

    row = FIRST_DATA_ROW
    for line in lines:
        _write_line(sheet, row, line, policy)
        row += 1

    last_row = max(row - 1, FIRST_DATA_ROW)
    _write_totals(sheet, FIRST_DATA_ROW, last_row, language)

    # Freeze the title, both header rows and the blank spacer, plus columns A-F.
    sheet.freeze_panes = f"{FIRST_GROUP_COLUMN}{FIRST_DATA_ROW}"
    _configure_print(sheet)
    return sheet


def build_strip_sheet(
    workbook: Workbook,
    title: str,
    period,
    lines,
    policy,
    language: str,
    branch_name: str | None = None,
) -> Worksheet:
    """Payslip strips: 2 header rows + data row + blank spacer, per employee.

    Printed and cut into individual strips, which is why each employee carries
    its own copy of the header.
    """
    sheet = workbook.create_sheet(title=title[:31])
    sheet.sheet_view.rightToLeft = language != "en"

    for column in COLUMNS:
        sheet.column_dimensions[column.letter].width = column.width

    row = 2
    for line in lines:
        header_top, header_sub, data_row = row, row + 1, row + 2

        earnings = GROUP_EARNINGS_EN if language == "en" else GROUP_EARNINGS_AR
        deductions = GROUP_DEDUCTIONS_EN if language == "en" else GROUP_DEDUCTIONS_AR
        sheet.merge_cells(f"{FIRST_GROUP_COLUMN}{header_top}:{LAST_EARNINGS_COLUMN}{header_top}")
        sheet.merge_cells(
            f"{FIRST_DEDUCTION_COLUMN}{header_top}:{LAST_DEDUCTION_COLUMN}{header_top}"
        )
        for letter, label in (
            (FIRST_GROUP_COLUMN, earnings),
            (FIRST_DEDUCTION_COLUMN, deductions),
        ):
            cell = sheet[f"{letter}{header_top}"]
            cell.value = label
            cell.font = HEADER_FONT
            cell.alignment = CENTER
            cell.fill = HEADER_FILL

        for column in COLUMNS:
            if column.group is None:
                sheet.merge_cells(f"{column.letter}{header_top}:{column.letter}{header_sub}")
                cell = sheet[f"{column.letter}{header_top}"]
            else:
                cell = sheet[f"{column.letter}{header_sub}"]
            cell.value = _header(column, language)
            cell.font = HEADER_FONT
            cell.alignment = CENTER
            cell.fill = HEADER_FILL
            cell.border = BORDER

        _write_line(sheet, data_row, line, policy)
        row += 4  # header, sub-header, data, blank spacer

    _configure_print(sheet)
    return sheet


def generate_workbook(period, lines, policy=None, language: str = "ar") -> bytes:
    """Build the whole report: master sheet, per-branch sheets, payslip strips."""
    formulas = PolicyFormulas(policy)
    workbook = Workbook()
    workbook.remove(workbook.active)

    ordered = sorted(lines, key=lambda line: (line.branch_snapshot or "", line.employee_code or ""))

    master_title = "Payroll report" if language == "en" else MASTER_SHEET_AR
    build_sheet(workbook, master_title, period, ordered, formulas, language)

    branches: dict[str, list] = {}
    for line in ordered:
        branches.setdefault(line.branch_snapshot or "-", []).append(line)

    for branch_name, branch_lines in branches.items():
        build_sheet(
            workbook,
            branch_name,
            period,
            branch_lines,
            formulas,
            language,
            branch_name=branch_name,
        )
        strip_title = (
            f"Payslips {branch_name}" if language == "en" else f"{STRIP_PREFIX_AR} {branch_name}"
        )
        build_strip_sheet(
            workbook,
            strip_title,
            period,
            branch_lines,
            formulas,
            language,
            branch_name=branch_name,
        )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def report_filename(company, period, extension: str = "xlsx") -> str:
    return f"{company.slug}-{period.year}-{period.month:02d}.{extension}"
