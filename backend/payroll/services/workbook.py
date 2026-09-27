"""Reader for the reference workbook `accounting_monthly.xlsx`.

Used by the parity test (to prove the engine matches) and by `seed_demo` (to
load the same 23 employees and February 2026 inputs into the database), so both
see exactly the same numbers.
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

import openpyxl

MASTER_SHEET = "تقرير الرواتب"
FIRST_DATA_ROW = 5
TOTALS_LABEL = "الاجماليات"

#: Excel column letter -> 1-based index, for the columns we read.
COLUMNS = {
    "code": 1,  # A
    "name": 2,  # B
    "branch": 3,  # C
    "job_title": 4,  # D
    "daily_hours": 5,  # E
    "base_salary": 6,  # F
    "work_days": 7,  # G
    "work_days_value": 8,  # H
    "overtime_hours": 9,  # I
    "overtime_value": 10,  # J
    "leave_allowance_days": 11,  # K
    "leave_allowance_value": 12,  # L
    "bonus": 13,  # M
    "other_earnings": 14,  # N
    "total_earnings": 15,  # O
    "late_hours": 16,  # P
    "late_value": 17,  # Q
    "advances": 18,  # R
    "carried_advance": 19,  # S
    "admin_penalty_days": 20,  # T
    "admin_penalty_value": 21,  # U
    "fingerprint_penalty_days": 22,  # V
    "fingerprint_penalty_value": 23,  # W
    "unexcused_absence_days": 24,  # X
    "unexcused_absence_value": 25,  # Y
    "sick_days": 26,  # Z
    "sick_value": 27,  # AA
    "insurable_salary": 28,  # AB
    "insurance_and_tax": 29,  # AC
    "deviations": 30,  # AD
    "shortage_custody": 31,  # AE
    "total_deductions": 32,  # AF
    "net_salary": 33,  # AG
    "notes": 34,  # AH
}

#: The engine's inputs, as read from the sheet.
INPUT_FIELDS = (
    "base_salary",
    "daily_hours",
    "work_days",
    "overtime_hours",
    "leave_allowance_days",
    "bonus",
    "other_earnings",
    "late_hours",
    "advances",
    "carried_advance",
    "admin_penalty_days",
    "fingerprint_penalty_days",
    "unexcused_absence_days",
    "sick_days",
    "insurable_salary",
    "deviations",
    "shortage_custody",
)

#: The calculated columns the engine must reproduce.
CALCULATED_FIELDS = (
    "work_days_value",
    "overtime_value",
    "leave_allowance_value",
    "total_earnings",
    "late_value",
    "admin_penalty_value",
    "fingerprint_penalty_value",
    "unexcused_absence_value",
    "sick_value",
    "insurance_and_tax",
    "total_deductions",
    "net_salary",
)


def _decimal(value) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    if isinstance(value, str):
        cleaned = value.strip().replace(",", "")
        if not cleaned:
            return Decimal("0")
        return Decimal(cleaned)
    return Decimal(str(value))


def _text(value) -> str:
    return "" if value is None else str(value).strip()


@dataclass(frozen=True)
class WorkbookRow:
    """One employee row from the master sheet, inputs and cached results."""

    row_number: int
    code: str
    name: str
    branch: str
    job_title: str
    inputs: dict[str, Decimal]
    expected: dict[str, Decimal]
    notes: str


@dataclass(frozen=True)
class WorkbookData:
    title: str
    rows: tuple[WorkbookRow, ...]
    totals: dict[str, Decimal]

    def by_code(self, code: str) -> WorkbookRow:
        for row in self.rows:
            if row.code == code:
                return row
        raise KeyError(code)


def load_reference(path: str | Path, sheet_name: str = MASTER_SHEET) -> WorkbookData:
    """Read employee rows and the totals row from the workbook's cached values."""
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook[sheet_name]
        grid = list(sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=34, values_only=True))
    finally:
        workbook.close()

    title = _text(grid[0][0]) if grid else ""

    rows: list[WorkbookRow] = []
    totals: dict[str, Decimal] = {}

    for index, raw in enumerate(grid, start=1):
        if index < FIRST_DATA_ROW:
            continue

        label = _text(raw[0])
        if label.startswith(TOTALS_LABEL):
            totals = {
                name: _decimal(raw[position - 1])
                for name, position in COLUMNS.items()
                if name not in {"code", "name", "branch", "job_title", "notes"}
            }
            break

        name = _text(raw[COLUMNS["name"] - 1])
        if not name:
            # Blank template row — the workbook keeps ~170 of them.
            continue

        rows.append(
            WorkbookRow(
                row_number=index,
                code=_text(raw[COLUMNS["code"] - 1]),
                name=name,
                branch=_text(raw[COLUMNS["branch"] - 1]),
                job_title=_text(raw[COLUMNS["job_title"] - 1]),
                inputs={f: _decimal(raw[COLUMNS[f] - 1]) for f in INPUT_FIELDS},
                expected={f: _decimal(raw[COLUMNS[f] - 1]) for f in CALCULATED_FIELDS},
                notes=_text(raw[COLUMNS["notes"] - 1]),
            )
        )

    return WorkbookData(title=title, rows=tuple(rows), totals=totals)
